#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
税后收入计算（表格版）- 2026年北京，累计预扣法
- 脚本本身不包含任何工资数据，数据全部由用户维护的 xlsx 文档提供。
- 输入：xlsx 第1行为表头，之后每行一个月，需包含以下列：
    月份 | 本月税前收入 | 累计税前收入 | 实际个税 | 累计专项附加扣除 | 每月五险一金
- 输出：计算结果作为新列追加在文档已有列之后（输入中已有的内容不再重复输出）。
- 用法：
    python 税后收入计算_表格版.py [xlsx路径]
    不传路径时默认使用与脚本同目录的“工资表.xlsx”。
"""
import sys
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

try:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
except ImportError:
    print("缺少 openpyxl，请先执行：pip install openpyxl")
    sys.exit(1)

# ---------------- 政策参数（2026年，可按政策变化自行更新） ----------------
MONTHLY_THRESHOLD = 5000  # 每月基本减除费用（起征点）
# (下限, 上限, 税率, 速算扣除数)，按“累计应纳税所得额”查档
TAX_BRACKETS = [
    (0, 36000, 0.03, 0),
    (36000, 144000, 0.10, 2520),
    (144000, 300000, 0.20, 16920),
    (300000, 420000, 0.25, 31920),
    (420000, 660000, 0.30, 52920),
    (660000, 960000, 0.35, 85920),
    (960000, float("inf"), 0.45, 181920),
]

# ---------------- xlsx 列约定 ----------------
REQUIRED_INPUTS = ["月份", "本月税前收入", "累计税前收入", "实际个税",
                   "累计专项附加扣除", "每月五险一金"]
RESULT_COLUMNS = ["累计减除费用", "累计五险一金", "累计收入差异", "累计应纳税所得额",
                  "适用税率", "速算扣除数", "累计应纳税额", "上月累计已纳税",
                  "本月应纳税额(计算)", "与工资单差异", "税后收入(计算)", "税后收入(实际)"]
MONEY_FORMAT = "0.00"
RATE_FORMAT = "0%"


def calc_cumulative_tax(taxable_income):
    """按累计应纳税所得额计算：(累计应纳税额, 适用税率, 速算扣除数)"""
    if taxable_income <= 0:
        return 0.0, 0.0, 0.0
    for low, high, rate, quick in TAX_BRACKETS:
        if low < taxable_income <= high:
            return taxable_income * rate - quick, rate, quick
    return 0.0, 0.0, 0.0


def read_number(cell, row_idx, col_name):
    """读取数值单元格，空值/非数值时报错退出。"""
    v = cell.value
    if v is None or (isinstance(v, str) and not v.strip()):
        print(f"错误：第{row_idx}行“{col_name}”为空，请补全后再运行。")
        sys.exit(1)
    try:
        return float(v)
    except (TypeError, ValueError):
        print(f"错误：第{row_idx}行“{col_name}”不是数值：{v!r}")
        sys.exit(1)


def main():
    default_path = Path(__file__).resolve().parent / "工资表.xlsx"
    path = Path(sys.argv[1]) if len(sys.argv) > 1 else default_path
    if not path.exists():
        print(f"错误：找不到文件 {path}")
        print("请先把工资表（xlsx）放到该路径，或以参数指定：python 税后收入计算_表格版.py <路径>")
        sys.exit(1)

    wb = load_workbook(path)
    ws = wb.active

    # 建立表头 -> 列号 映射（按第1行，去空格匹配）
    header_map = {}
    for col in range(1, ws.max_column + 1):
        v = ws.cell(row=1, column=col).value
        if v is not None and str(v).strip():
            header_map[str(v).strip()] = col

    missing = [h for h in REQUIRED_INPUTS if h not in header_map]
    if missing:
        print("错误：xlsx 缺少必需列：" + "、".join(missing))
        print(f"当前第1行表头：{list(header_map.keys())}")
        sys.exit(1)

    # 结果列定位：已存在同名表头则原地覆盖（支持重算），否则在现有列之后追加
    next_col = ws.max_column + 1
    result_cols = {}
    for name in RESULT_COLUMNS:
        if name in header_map:
            result_cols[name] = header_map[name]
        else:
            result_cols[name] = next_col
            ws.cell(row=1, column=next_col, value=name)
            next_col += 1

    def put(row_idx, name, value, fmt=MONEY_FORMAT):
        cell = ws.cell(row=row_idx, column=result_cols[name], value=round(value, 2))
        cell.number_format = fmt

    # 逐行计算
    running_income = 0.0
    cum_tax_paid = 0.0     # 之前月份累计“已按本方法计算”的税额
    rows_meta = []         # (月份, 与工资单差异) 供控制台提示
    total_calc_tax = total_actual_tax = total_insurance = 0.0
    n_months = 0

    r = 2
    while True:
        month_cell = ws.cell(row=r, column=header_map["月份"]).value
        if month_cell is None or (isinstance(month_cell, str) and not month_cell.strip()):
            break
        month = int(float(month_cell))
        income = read_number(ws.cell(row=r, column=header_map["本月税前收入"]), r, "本月税前收入")
        cum_income = read_number(ws.cell(row=r, column=header_map["累计税前收入"]), r, "累计税前收入")
        actual_tax = read_number(ws.cell(row=r, column=header_map["实际个税"]), r, "实际个税")
        cum_special = read_number(ws.cell(row=r, column=header_map["累计专项附加扣除"]), r, "累计专项附加扣除")
        insurance = read_number(ws.cell(row=r, column=header_map["每月五险一金"]), r, "每月五险一金")

        running_income += income
        cum_threshold = MONTHLY_THRESHOLD * month
        cum_insurance = insurance * month
        taxable = cum_income - cum_threshold - cum_insurance - cum_special
        cum_tax, rate, quick = calc_cumulative_tax(taxable)
        current_tax = max(0.0, cum_tax - cum_tax_paid)
        diff = current_tax - actual_tax

        put(r, "累计减除费用", cum_threshold)
        put(r, "累计五险一金", cum_insurance)
        put(r, "累计收入差异", cum_income - running_income)
        put(r, "累计应纳税所得额", taxable if taxable > 0 else 0.0)
        put(r, "适用税率", rate, RATE_FORMAT)
        put(r, "速算扣除数", quick)
        put(r, "累计应纳税额", cum_tax)
        put(r, "上月累计已纳税", cum_tax_paid)
        put(r, "本月应纳税额(计算)", current_tax)
        put(r, "与工资单差异", diff)
        put(r, "税后收入(计算)", income - insurance - current_tax)
        put(r, "税后收入(实际)", income - insurance - actual_tax)

        rows_meta.append((month, diff))
        cum_tax_paid = cum_tax
        total_calc_tax += current_tax
        total_actual_tax += actual_tax
        total_insurance += insurance
        n_months += 1
        r += 1

    if n_months == 0:
        print("错误：第2行起未找到任何数据行（“月份”列为空）。")
        sys.exit(1)

    wb.save(path)

    # ---------------- 控制台简报（不重复输出文档已有内容） ----------------
    first_new = min(result_cols.values())
    last_new = max(result_cols.values())
    print(f"计算完成：共 {n_months} 个月，结果已写入《{path.name}》"
          f"第 {get_column_letter(first_new)}–{get_column_letter(last_new)} 列。")
    bad = [(m, d) for m, d in rows_meta if abs(d) > 0.005]
    if bad:
        print("注意：以下月份“本月应纳税额(计算)”与工资单不一致，请核对输入：")
        for m, d in bad:
            print(f"  {m}月 差异 {d:+.2f} 元")
    else:
        print("校验：各月计算税额与工资单实际个税一致。")
    print(f"累计：个税(计算) {total_calc_tax:.2f} 元 | 五险一金 {total_insurance:.2f} 元")


if __name__ == "__main__":
    main()